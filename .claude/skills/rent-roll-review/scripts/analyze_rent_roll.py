#!/usr/bin/env python3
"""
Analyze a Yardi Rent Roll export (.xlsx) and flag items needing attention.

Usage:
    python3 analyze_rent_roll.py <path-to-rent-roll.xlsx> [--as-of MM/DD/YYYY]

Produces:
  - Portfolio snapshot (units, occupancy, monthly rent)
  - Per-property summary
  - Exception flags: delinquencies, large credit balances, expiring/expired
    leases, move-outs, vacancies, $0-deposit occupied units, $0-rent occupied
    units, negative misc charges, and future-applicant pipeline.

The parser is format-driven (not hardcoded to specific property names). It keys
off the standard Yardi layout: section headers ("Current/Notice/Vacant Tenants",
"Future Tenants/Applicants") and per-property "Total" rows whose property name
sits in the Tenant-Name column.
"""
import sys
import argparse
from datetime import datetime, timedelta

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required. Install with: pip install openpyxl")

# Standard Yardi Rent Roll column layout (1-indexed)
COL = dict(unit=1, sqft=2, name=3, rent=4, rent_psf=5, dep=6, other_dep=7,
           misc=8, misc_psf=9, movein=10, leaseexp=11, moveout=12, balance=13)

SEC_CURRENT = "Current/Notice/Vacant Tenants"
SEC_FUTURE = "Future Tenants/Applicants"

# Thresholds (tune as needed)
LARGE_CREDIT = -2000.0      # credit balances more negative than this get flagged
EXPIRY_HORIZON_DAYS = 60    # leases expiring within this window are flagged


def parse_date(v):
    if isinstance(v, datetime):
        return v
    if not v:
        return None
    for fmt in ("%m/%d/%Y", "%Y-%m-%d", "%m/%d/%y"):
        try:
            return datetime.strptime(str(v).strip(), fmt)
        except ValueError:
            continue
    return None


def num(v):
    return v if isinstance(v, (int, float)) else None


def is_vacant(name):
    return str(name).strip().upper().startswith("VACANT")


def load_rows(path):
    """Return (rows, meta). Each row is a dict tagged with property + section."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]

    def cell(r, key):
        return ws.cell(row=r, column=COL[key]).value

    meta = {}
    for r in range(1, 6):
        a = ws.cell(row=r, column=1).value
        if isinstance(a, str) and "As Of" in a:
            meta["as_of"] = a.split("=", 1)[-1].strip()
        if isinstance(a, str) and a.strip().startswith("Property"):
            meta["properties"] = a.split("=", 1)[-1].strip()

    pending = []       # detail rows awaiting their property's Total row
    section = None
    rows = []
    for r in range(1, ws.max_row + 1):
        a = ws.cell(row=r, column=1).value
        name = cell(r, "name")

        if a in (SEC_CURRENT, SEC_FUTURE):
            section = a
            continue

        # Per-property Total row: closes out the current block.
        if a == "Total" and name and str(name).strip().lower() != "all properties":
            prop = str(name).strip()
            for rec in pending:
                rec["property"] = prop
                rows.append(rec)
            pending = []
            section = None
            continue

        # Stop once we hit the All-Properties total / Summary Groups footer.
        if a == "Total" and str(name).strip().lower() == "all properties":
            break
        if isinstance(a, str) and a.strip() == "Summary Groups":
            break

        # Detail row: has a unit id and a tenant name.
        if a is not None and name is not None and section is not None:
            pending.append(dict(
                section=section,
                unit=str(a).strip(),
                sqft=num(cell(r, "sqft")),
                name=str(name).strip(),
                rent=num(cell(r, "rent")),
                dep=num(cell(r, "dep")),
                misc=num(cell(r, "misc")),
                movein=cell(r, "movein"),
                leaseexp=cell(r, "leaseexp"),
                moveout=cell(r, "moveout"),
                balance=num(cell(r, "balance")),
            ))

    return rows, meta


def money(x):
    return f"${x:,.2f}"


def analyze(path, as_of_override=None):
    rows, meta = load_rows(path)
    as_of = parse_date(as_of_override or meta.get("as_of")) or datetime.now()
    cutoff = as_of + timedelta(days=EXPIRY_HORIZON_DAYS)

    props = {}
    for x in rows:
        props.setdefault(x["property"], []).append(x)

    out = []
    out.append("=" * 70)
    out.append("RENT ROLL REVIEW")
    out.append(f"As Of: {as_of.strftime('%m/%d/%Y')}")
    if meta.get("properties"):
        out.append(f"Properties: {meta['properties']}")
    out.append("=" * 70)

    # Portfolio rollup
    all_cur = [x for x in rows if x["section"] == SEC_CURRENT]
    occ = [x for x in all_cur if not is_vacant(x["name"])]
    vac = [x for x in all_cur if is_vacant(x["name"])]
    total_rent = sum(x["rent"] or 0 for x in occ)
    occ_pct = 100 * len(occ) / len(all_cur) if all_cur else 0
    owed = [x for x in all_cur if (x["balance"] or 0) > 0]
    total_owed = sum(x["balance"] for x in owed)
    out.append("\nPORTFOLIO SNAPSHOT")
    out.append(f"  Units: {len(all_cur)}   Occupied: {len(occ)}   "
               f"Vacant: {len(vac)}   Occupancy: {occ_pct:.1f}%")
    out.append(f"  Monthly rent (occupied): {money(total_rent)}")
    out.append(f"  Outstanding owed: {money(total_owed)} across {len(owed)} tenant(s)")

    # Per-property
    out.append("\n" + "-" * 70)
    out.append("PER-PROPERTY SUMMARY")
    for prop in props:
        cur = [x for x in props[prop] if x["section"] == SEC_CURRENT]
        fut = [x for x in props[prop] if x["section"] == SEC_FUTURE]
        po = [x for x in cur if not is_vacant(x["name"])]
        pv = [x for x in cur if is_vacant(x["name"])]
        prent = sum(x["rent"] or 0 for x in po)
        powed = [x for x in cur if (x["balance"] or 0) > 0]
        pcred = [x for x in cur if (x["balance"] or 0) < 0]
        pmo = [x for x in po if x["moveout"]]
        pct = 100 * len(po) / len(cur) if cur else 0
        out.append(f"\n  {prop}")
        out.append(f"    Units {len(cur)} | Occupied {len(po)} | Vacant {len(pv)} "
                   f"| Occupancy {pct:.1f}%")
        out.append(f"    Monthly rent: {money(prent)}")
        out.append(f"    Owed: {money(sum(x['balance'] for x in powed))} ({len(powed)}) | "
                   f"Credits: {money(sum(x['balance'] for x in pcred))} ({len(pcred)})")
        out.append(f"    Move-outs on notice: {len(pmo)} | Future applicants: {len(fut)}")
        if pv:
            out.append("    Vacant: " + ", ".join(x["unit"] for x in pv))

    # Exception flags
    def line(x, extra=""):
        return (f"    {x['property'][:20]:20} U{x['unit']:>8} "
                f"{x['name'][:28]:28} {extra}")

    out.append("\n" + "-" * 70)
    out.append("FLAGS NEEDING ATTENTION")

    out.append(f"\n  [1] DELINQUENCIES — {money(total_owed)} owed:")
    for x in sorted(owed, key=lambda r: -r["balance"]):
        note = ""
        mo = parse_date(x["moveout"])
        if mo:
            note = "<<< already moved/moving out" if mo <= cutoff else ""
        out.append(line(x, f"{money(x['balance']):>12}  {note}"))

    big_cred = [x for x in all_cur if (x["balance"] or 0) <= LARGE_CREDIT]
    out.append(f"\n  [2] LARGE CREDIT BALANCES (<= {money(LARGE_CREDIT)}) — verify postings:")
    for x in sorted(big_cred, key=lambda r: r["balance"]):
        out.append(line(x, f"{money(x['balance']):>12}"))

    expiring = []
    for x in occ:
        d = parse_date(x["leaseexp"])
        if d and d <= cutoff:
            expiring.append((d, x))
    out.append(f"\n  [3] LEASES EXPIRED / EXPIRING within {EXPIRY_HORIZON_DAYS} days:")
    for d, x in sorted(expiring, key=lambda t: t[0]):
        tag = "*** EXPIRED" if d < as_of else ""
        mo = "moveout " + str(x["moveout"]).split()[0] if x["moveout"] else "no notice"
        out.append(line(x, f"exp {d.strftime('%m/%d/%Y')}  {mo:18} {tag}"))

    moveouts = [x for x in occ if x["moveout"]]
    out.append(f"\n  [4] MOVE-OUTS ON NOTICE — {len(moveouts)}:")
    for x in sorted(moveouts, key=lambda r: parse_date(r["moveout"]) or as_of):
        bal = f"owes {money(x['balance'])}" if (x["balance"] or 0) > 0 else ""
        out.append(line(x, f"out {str(x['moveout']).split()[0]:12} {bal}"))

    zero_dep = [x for x in occ if x["dep"] == 0]
    out.append(f"\n  [5] OCCUPIED UNITS WITH $0 SECURITY DEPOSIT — {len(zero_dep)}:")
    for x in zero_dep:
        out.append(line(x, f"rent {money(x['rent'] or 0)}"))

    zero_rent = [x for x in occ if not x["rent"]]
    out.append(f"\n  [6] OCCUPIED UNITS WITH $0 RENT — {len(zero_rent)}:")
    for x in zero_rent:
        out.append(line(x))

    neg_misc = [x for x in occ if (x["misc"] or 0) < 0]
    out.append(f"\n  [7] NEGATIVE MISC CHARGES (concessions/credits) — {len(neg_misc)}:")
    for x in neg_misc:
        out.append(line(x, f"misc {money(x['misc'])}"))

    print("\n".join(out))


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Analyze a Yardi rent roll .xlsx")
    ap.add_argument("path", help="Path to the rent roll .xlsx file")
    ap.add_argument("--as-of", dest="as_of", default=None,
                    help="Override the As-Of date (MM/DD/YYYY)")
    args = ap.parse_args()
    analyze(args.path, args.as_of)
