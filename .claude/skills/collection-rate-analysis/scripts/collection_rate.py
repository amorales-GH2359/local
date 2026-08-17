#!/usr/bin/env python3
"""Collection-rate analysis for CS3 "Rent Roll" AR Excel exports.

Usage:
    python3 collection_rate.py FILE1.xlsx [FILE2.xlsx ...]

For each property block it finds, computes:
    Billed          = Actual Rent total + Misc total   (current tenants only)
    Uncollected AR  = sum of positive Balance values    (current tenants only)
    Collection Rate = (Billed - Uncollected AR) / Billed * 100

Rules: only positive balances count as uncollected AR (negatives are credits and
are ignored); future tenants/applicants are excluded; Residential and Commercial
sections are reported separately. Rent/Misc are summed directly from the
current-tenant rows (not the section Total row) so future-tenant rent never
inflates the billed base.
"""
import sys
import openpyxl

RENT_COL, MISC_COL, BAL_COL, UNIT_COL, NAME_COL = 4, 8, 13, 1, 3


def num(v):
    return v if isinstance(v, (int, float)) else None


def analyze(path):
    ws = openpyxl.load_workbook(path, data_only=True)["Report1"]

    # Sanity check: this skill only handles the "Rent Roll" export.
    a1 = str(ws.cell(1, 1).value or "")
    if "Rent Roll" not in a1:
        raise SystemExit(
            f"{path}: A1 is {a1!r}, not a 'Rent Roll' report. "
            "This script does not handle other exports (e.g. 'Receivable Summary')."
        )

    prop = next((ws.cell(r, 1).value for r in range(1, 6)
                 if str(ws.cell(r, 1).value or "").startswith("Property")), path)
    asof = next((ws.cell(r, 1).value for r in range(1, 6)
                 if str(ws.cell(r, 1).value or "").startswith("As Of")), "")

    sections, cur, in_future = [], None, False
    for r in range(1, ws.max_row + 1):
        a = str(ws.cell(r, UNIT_COL).value or "").strip()
        if a == "Current/Notice/Vacant Tenants":
            cur, in_future = {"pos": 0.0, "items": [], "rent": 0.0, "misc": 0.0}, False
        elif a == "Future Tenants/Applicants":
            in_future = True
        elif a == "Total" and cur is not None:
            cur["name"] = ws.cell(r, NAME_COL).value
            sections.append(cur)
            cur, in_future = None, False
        elif cur is not None and not in_future:
            cur["rent"] += num(ws.cell(r, RENT_COL).value) or 0.0
            cur["misc"] += num(ws.cell(r, MISC_COL).value) or 0.0
            bal = num(ws.cell(r, BAL_COL).value)
            if bal and bal > 0:
                cur["pos"] += bal
                cur["items"].append((a, ws.cell(r, NAME_COL).value, bal))

    return prop, asof, [s for s in sections if "name" in s]


def main(paths):
    for path in paths:
        prop, asof, sections = analyze(path)
        print("=" * 78)
        print(prop, f"  ({asof})  [{path}]")
        for s in sections:
            billed = s["rent"] + s["misc"]
            collected = billed - s["pos"]
            rate = collected / billed * 100 if billed else float("nan")
            print(f"\n  {s['name']}")
            print(f"    Actual Rent : {s['rent']:>12,.2f}")
            print(f"    Misc        : {s['misc']:>12,.2f}")
            print(f"    Billed      : {billed:>12,.2f}")
            print(f"    Positive AR : {s['pos']:>12,.2f}")
            print(f"    Collected   : {collected:>12,.2f}")
            print(f"    RATE        : {rate:>11.2f}%")
            if s["items"]:
                print("    Uncollected (positive AR) tenants:")
                for unit, name, bal in s["items"]:
                    print(f"      {unit:<10} {str(name)[:32]:<32} {bal:>10,.2f}")
        print()
    print("NOTE: Rent/Misc are summed from current-tenant rows only; future")
    print("tenants/applicants and negative (credit) balances are excluded.")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    main(sys.argv[1:])
