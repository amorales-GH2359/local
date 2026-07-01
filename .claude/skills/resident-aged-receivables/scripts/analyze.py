#!/usr/bin/env python3
"""Parse Resident Aged Receivables workbooks and print rows, totals, and
auto-flagged anomalies for each property.

Usage:
    python3 analyze.py FILE1.xlsx [FILE2.xlsx ...]

Requires openpyxl (`pip install openpyxl`). Reads with data_only=True so
formula cells return their last computed value.
"""
import re
import sys

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required: pip install openpyxl")

# Strips the "MM/DD/YYYY HH:MM AM/PM someone@domain:" prefix from a note so we
# can compare the message body across residents (timestamps often differ even
# when the text was copy/pasted).
NOTE_PREFIX = re.compile(r"^\s*\d{1,2}/\d{1,2}/\d{2,4}.*?@\S+?:")


def note_body(text):
    return NOTE_PREFIX.sub("", str(text)).strip().lower()

# Header labels expected in row 6 of the data sheet, in order.
COLS = [
    "Bldg-Unit", "Resident", "Lease Status", "Unallocated Charges / Credits",
    "0-30 Days", "31-60 Days", "61-90 Days", "90+ Days", "Pre-Payments",
    "Balance", "Last Delinquency Note",
]
# Lease statuses that normally should NOT carry an outstanding balance.
SUSPECT_STATUSES = {"cancelled", "applicant", "future"}


def money(v):
    return f"${v:,.2f}" if isinstance(v, (int, float)) else str(v)


def find_header_row(ws):
    for r in range(1, min(ws.max_row, 15) + 1):
        if ws.cell(r, 1).value == "Bldg-Unit":
            return r
    return 6


def analyze_sheet(ws, params):
    hr = find_header_row(ws)
    rows, total_row = [], None
    for r in range(hr + 1, ws.max_row + 1):
        unit = ws.cell(r, 1).value
        label = ws.cell(r, 3).value  # "<Property> Total:" lands in col C
        if unit is None and isinstance(label, str) and "Total" in label:
            total_row = {COLS[i]: ws.cell(r, i + 1).value for i in range(len(COLS))}
            continue
        if unit is None:
            continue
        rows.append({COLS[i]: ws.cell(r, i + 1).value for i in range(len(COLS))})
    return rows, total_row


def num(x):
    return x if isinstance(x, (int, float)) else 0.0


def report(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    data_ws = wb.worksheets[0]
    params = {}
    if "Report Parameters" in wb.sheetnames:
        pw = wb["Report Parameters"]
        for r in range(1, pw.max_row + 1):
            k, v = pw.cell(r, 1).value, pw.cell(r, 2).value
            if k:
                params[str(k).strip()] = v

    prop = data_ws.cell(3, 1).value or data_ws.title
    post_month = data_ws.cell(4, 1).value
    rows, total_row = analyze_sheet(data_ws, params)

    print("#" * 100)
    print(f"PROPERTY: {prop}   post month: {post_month}")
    print(f"  data as of: {params.get('data as of') or params.get(' data as of')}")
    print(f"  residents listed: {len(rows)}")
    print("-" * 100)

    for row in rows:
        print(f"  {row['Bldg-Unit']:<16} {str(row['Resident'])[:34]:<34} "
              f"[{row['Lease Status']}]")
        print(f"      Balance={money(row['Balance'])}  0-30={money(row['0-30 Days'])}  "
              f"31-60={money(row['31-60 Days'])}  61-90={money(row['61-90 Days'])}  "
              f"90+={money(row['90+ Days'])}  Pre-Pay={money(row['Pre-Payments'])}")
        if row["Last Delinquency Note"]:
            print(f"      note: {row['Last Delinquency Note']}")

    # ---- recompute totals and compare to the sheet's total row ----
    keys = ["Unallocated Charges / Credits", "0-30 Days", "31-60 Days",
            "61-90 Days", "90+ Days", "Pre-Payments", "Balance"]
    calc = {k: sum(num(r[k]) for r in rows) for k in keys}
    print("-" * 100)
    print("  COMPUTED TOTALS:", "  ".join(f"{k}={money(v)}" for k, v in calc.items()))
    if total_row:
        for k in keys:
            sheet_v = num(total_row.get(k))
            if abs(sheet_v - calc[k]) > 0.01:
                print(f"  !! TOTAL MISMATCH on {k}: sheet={money(sheet_v)} "
                      f"computed={money(calc[k])}")

    # ---- anomaly flags ----
    print("-" * 100)
    print("  FLAGS:")
    flagged = False

    def flag(msg):
        nonlocal flagged
        flagged = True
        print(f"    - {msg}")

    # credit balances
    for r in rows:
        if num(r["Balance"]) < 0:
            flag(f"CREDIT balance {money(r['Balance'])} — {r['Bldg-Unit']} "
                 f"{r['Resident']} (prepayment, not past-due)")
    # aged debt
    for r in rows:
        if num(r["90+ Days"]) > 0:
            flag(f"90+ DAYS {money(r['90+ Days'])} — {r['Bldg-Unit']} {r['Resident']}")
    for r in rows:
        if num(r["61-90 Days"]) > 0:
            flag(f"61-90 days {money(r['61-90 Days'])} — {r['Bldg-Unit']} {r['Resident']}")
    # suspect lease status with a positive balance
    for r in rows:
        st = str(r["Lease Status"] or "").strip().lower()
        if any(s in st for s in SUSPECT_STATUSES) and num(r["Balance"]) > 0:
            flag(f"Lease status '{r['Lease Status']}' carries balance "
                 f"{money(r['Balance'])} — {r['Bldg-Unit']} {r['Resident']}")
    # duplicate unit numbers
    seen = {}
    for r in rows:
        seen.setdefault(r["Bldg-Unit"], []).append(r["Resident"])
    for unit, residents in seen.items():
        if len(residents) > 1:
            flag(f"DUPLICATE unit {unit}: {', '.join(map(str, residents))}")
    # large balance with no note (threshold: top-quartile-ish, >= $1000)
    for r in rows:
        if num(r["Balance"]) >= 1000 and not r["Last Delinquency Note"]:
            flag(f"Large balance {money(r['Balance'])} with NO note — "
                 f"{r['Bldg-Unit']} {r['Resident']}")
    # duplicated note text across residents
    notes = {}
    for r in rows:
        n = r["Last Delinquency Note"]
        if n:
            body = note_body(n)
            if body:
                notes.setdefault(body, []).append(f"{r['Bldg-Unit']} {r['Resident']}")
    for body, who in notes.items():
        if len(who) > 1:
            flag(f"IDENTICAL note body on {len(who)} residents ({'; '.join(who)}) "
                 f"— possible copy/paste/stale note")
    # entity/corporate accounts
    for r in rows:
        name = str(r["Resident"] or "")
        if name.startswith("LLC") or "LLC," in name or "Workforce" in name:
            flag(f"Entity/corporate account — {r['Bldg-Unit']} {name}")

    if not flagged:
        print("    (none)")
    print()


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    for path in sys.argv[1:]:
        report(path)


if __name__ == "__main__":
    main()
